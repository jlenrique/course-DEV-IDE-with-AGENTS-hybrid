"""Story 39.1 — glossary downstream integration: _act intake + composed surfaces.

Rows h (MD association), i (DOCX association via python-docx extraction),
j (citation resolvability into the cited-entries block) + the AC-A9 dedupe,
plus the A-3 intake pins: status-dependent authority served from ONE 07W.3
disk read, hoisted above the glossary intake.

Live-shape fixtures only: the run rig carries the frozen a940c5eb workbook
brief + the real 1-row Ask-A pool (via the 37-2b helper substrate).
"""

from __future__ import annotations

import json
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pytest
import yaml
from docx import Document

from app.marcus.lesson_plan.deep_dive_enrichment import (
    DeepDiveEnrichmentExecutionReceiptV1,
    build_workbook_review_contribution,
    compose_deep_dive_enrichment,
)
from app.marcus.lesson_plan.prework_artifact import read_workbook_brief
from app.marcus.lesson_plan.workbook_producer import WorkbookProducer
from app.models.runtime.production_envelope import (
    ProductionEnvelope,
    SpecialistContribution,
)
from app.models.runtime.production_trial_envelope import ProductionTrialEnvelope
from app.specialists.workbook_producer import _act as wb_act
from scripts.utilities import marcus_spoc_live_test_runner as runner
from tests.helpers.deep_dive_enrichment_37_2b import (
    enriched_candidate,
    install_brief,
    live_skeleton_binding,
    make_request,
    mutated_pool_entry,
)

from ._run_fixture import collateral_present, section

LIVE_TERMS = tuple(marker.term for marker in live_skeleton_binding().bold_terms)

# The corpus carries an in-source numeral so the G1 audit's zero-denominator
# guard does not fire (per the B5 precedent); deep-dive/glossary figure tokens
# are cleared via the declared render supplements.
_CORPUS = "Administrative waste is roughly 25% of total spend.\n"
_SEGMENTS = [
    {
        "segment_id": "seg-01",
        "id": "seg-01",
        "slide_id": "slide-01",
        "narration_text": "Innovating inside the hospital means navigating legacy systems.",
    }
]
_COLLATERAL = collateral_present(
    [
        section(
            "sec-u01",
            "u01",
            title="Reading the structural shift",
            deferred_depth="Read-channel derivation of the structural-shift objective.",
        )
    ]
)

_BOLD_RE = re.compile(r"\*\*([^*]+)\*\*")


def _receipt(request) -> DeepDiveEnrichmentExecutionReceiptV1:
    return DeepDiveEnrichmentExecutionReceiptV1(
        mode="live",
        calls=1,
        idempotency_key="sha256:" + "2" * 64,
        request_digest=request.request_digest,
        pool_packet_digest=request.pool_packet_digest,
        model="gpt-5",
        model_config_digest="sha256:" + "3" * 64,
        cost_usd=0.05,
    )


def _enriched_contribution_payload(*, extra_pool_row: bool = False) -> dict[str, Any]:
    kwargs: dict[str, Any] = {}
    if extra_pool_row:
        from tests.helpers.deep_dive_enrichment_37_2b import live_pool_entry

        kwargs["pool_rows"] = (
            live_pool_entry(),
            mutated_pool_entry(citation_id="ask-a-cite-002"),
        )
    request = make_request(**kwargs)
    result = compose_deep_dive_enrichment(request, lambda _: enriched_candidate(request))
    contribution = build_workbook_review_contribution(result, _receipt(request))
    return json.loads(contribution.model_dump_json())


def _make_run_dir(root: Path, contribution_payload: dict[str, Any] | None) -> Path:
    run_dir = root / "run"
    (run_dir / "exports").mkdir(parents=True)
    (run_dir / "bundle").mkdir()
    (run_dir / "exports" / "segment-manifest-storyboard-b.yaml").write_text(
        yaml.safe_dump({"segments": _SEGMENTS}, sort_keys=False), encoding="utf-8"
    )
    (run_dir / "bundle" / "extracted.md").write_text(_CORPUS, encoding="utf-8")
    install_brief(run_dir)
    trial_id = uuid.uuid4()
    envelope = ProductionEnvelope(trial_id=trial_id)
    envelope.add_contribution(
        SpecialistContribution.from_output(
            specialist_id="irene_pass1",
            output={
                "lesson_plan": {
                    "lesson_summary": "glossary downstream 39.1",
                    "plan_units": [],
                    "collateral": _COLLATERAL,
                }
            },
            model_used="fixture-irene",
            node_id="03",
        )
    )
    if contribution_payload is not None:
        envelope.add_contribution(
            SpecialistContribution.from_output(
                specialist_id="workbook_review",
                output=contribution_payload,
                model_used="gpt-5",
                node_id="07W.3",
            )
        )
    trial = ProductionTrialEnvelope(
        trial_id=trial_id,
        preset="production",
        corpus_path="course-content/courses/fixture-lesson",
        operator_id="fixture-operator",
        started_at=datetime.now(UTC),
        status="in-flight",
        production_clone_launch_evidence=False,
        production_envelope=envelope,
    )
    (run_dir / "run.json").write_text(trial.model_dump_json(), encoding="utf-8")
    return run_dir


def _build_inputs(run_dir: Path) -> Any:
    brief = read_workbook_brief(run_dir)
    inputs = wb_act.build_workbook_inputs(
        run_dir, run_id="glos39", validated_brief=brief
    )
    assert inputs is not None
    return inputs


def _produce(run_dir: Path, inputs: Any) -> tuple[str, Path]:
    producer = WorkbookProducer(output_root=str(run_dir / "exports" / "workbooks"))
    sidecar = producer.produce(
        inputs.plan_unit,
        inputs.context,
        workbook_title=inputs.workbook_title,
        spec=inputs.spec,
        segments=inputs.segments,
        source_text=inputs.source_text,
        citations=inputs.citations,
        source_ref_manifest=inputs.source_ref_manifest,
        vo_script_text=inputs.vo_script_text,
        learning_objectives=inputs.learning_objectives,
        answer_keys=inputs.answer_keys,
        further_reading=inputs.further_reading,
        research_entries=inputs.research_entries,
        research_empty_reason=inputs.research_empty_reason,
        research_omitted_note=inputs.research_omitted_note,
        glossary_articles=inputs.glossary_articles,
        glossary_empty_reason=inputs.glossary_empty_reason,
        glossary=inputs.glossary,
        research_trends=inputs.research_trends,
        research_supplements=inputs.research_supplements,
        lo_overlay_loss=inputs.lo_overlay_loss,
        pre_work=inputs.pre_work,
        encounter_mode=inputs.encounter_mode,
        render_profile=inputs.render_profile,
        workbook_brief_receipt=inputs.workbook_brief_receipt,
        deep_dive_review=inputs.deep_dive_review,
    )
    repo_root = Path(wb_act.REPO_ROOT)
    md_path = repo_root / sidecar.markdown_path
    markdown = md_path.read_text(encoding="utf-8")
    return markdown, repo_root / sidecar.docx_path


def _section_body(md: str, heading: str) -> str:
    marker = f"\n## {heading}\n"
    start = md.index(marker) + len(marker)
    rest = md[start:]
    nxt = rest.find("\n## ")
    return rest if nxt == -1 else rest[:nxt]


# ---------------------------------------------------------------------------
# _act intake — status-dependent authority off ONE hoisted 07W.3 read (A-3)
# ---------------------------------------------------------------------------


def test_intake_enriched_authority_and_manifest_join(tmp_path: Path) -> None:
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    inputs = _build_inputs(run_dir)
    glossary = inputs.glossary
    assert glossary is not None and glossary.authority == "enriched"
    assert tuple(e.term for e in glossary.entries) == LIVE_TERMS
    assert glossary.covered_count == 1
    # Covered articles ride the legacy-compatible fields for G2/G1 seams.
    assert inputs.glossary_articles == glossary.covered_articles
    article = inputs.glossary_articles[0]
    # Manifest join carries the pool row's OWN hash (no R19 disagreement).
    assert inputs.source_ref_manifest[article.source_ref] == article.source_hash


def test_intake_absent_contribution_yields_authority_absent(tmp_path: Path) -> None:
    run_dir = _make_run_dir(tmp_path, None)
    inputs = _build_inputs(run_dir)
    assert inputs.glossary is not None
    assert inputs.glossary.authority == "absent"
    assert inputs.glossary_articles == ()
    assert "bold-term authority absent" in (inputs.glossary_empty_reason or "")


def test_intake_single_disk_read_serves_glossary_and_deep_dive(
    tmp_path: Path, monkeypatch
) -> None:
    """A-3 hoist: the 07W.3 contribution is read from disk exactly ONCE."""
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    calls = {"n": 0}
    real = wb_act.load_workbook_review_contribution

    def counting(load_dir: Path):
        calls["n"] += 1
        return real(load_dir)

    monkeypatch.setattr(wb_act, "load_workbook_review_contribution", counting)
    inputs = _build_inputs(run_dir)
    assert calls["n"] == 1
    assert inputs.deep_dive_review is not None
    assert inputs.glossary is not None and inputs.glossary.authority == "enriched"


# ---------------------------------------------------------------------------
# Rows h + i + j on the PRODUCED deliverable (MD + DOCX + references)
# ---------------------------------------------------------------------------


def test_row_h_md_bold_set_equals_glossary_headword_set(tmp_path: Path) -> None:
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    markdown, _ = _produce(run_dir, _build_inputs(run_dir))
    deep_dive = _section_body(markdown, "Deep Dive")
    glossary = _section_body(markdown, "Research Glossary")
    bold_terms = set(_BOLD_RE.findall(deep_dive))
    headwords = {
        line[4:].strip() for line in glossary.splitlines() if line.startswith("### ")
    }
    assert bold_terms == set(LIVE_TERMS)
    assert headwords == bold_terms  # missing entry = FAIL; orphan entry = FAIL


def test_row_i_docx_bold_runs_match_glossary_heading3(tmp_path: Path) -> None:
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    _, docx_path = _produce(run_dir, _build_inputs(run_dir))
    document = Document(str(docx_path))
    region: str | None = None
    deep_dive_bolds: set[str] = set()
    glossary_h3: set[str] = set()
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if style == "Heading 2":
            region = paragraph.text
            continue
        if region == "Deep Dive":
            for run in paragraph.runs:
                if run.bold and run.text.strip():
                    deep_dive_bolds.add(run.text)
        elif region == "Research Glossary" and style == "Heading 3":
            glossary_h3.add(paragraph.text)
    assert deep_dive_bolds == set(LIVE_TERMS)
    assert glossary_h3 == deep_dive_bolds


def test_row_j_covered_citation_resolves_once_in_cited_entries(tmp_path: Path) -> None:
    """Dedupe (AC-A9): ask-a-cite-001 is BOTH prose-used and glossary-covering —
    it appears exactly ONCE in the cited-entries block."""
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    markdown, _ = _produce(run_dir, _build_inputs(run_dir))
    references = _section_body(markdown, "References")
    assert "#### Deep Dive cited entries (Ask-A, DOI)" in references
    assert references.count("citation_id: `ask-a-cite-001`") == 1


def test_row_j_glossary_only_citation_appended_and_resolvable(tmp_path: Path) -> None:
    """A pool row that covers a term but is UNUSED by prose is appended into
    the cited-entries block by the AC-A9 seam (probe row-j dependency)."""
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload(extra_pool_row=True))
    markdown, _ = _produce(run_dir, _build_inputs(run_dir))
    references = _section_body(markdown, "References")
    assert references.count("citation_id: `ask-a-cite-001`") == 1
    assert references.count("citation_id: `ask-a-cite-002`") == 1
    # And the full deliverable bar (deep-dive + glossary clauses) accepts the
    # real produced artifact end-to-end.
    runner._assert_completed_workbook_deliverable(uuid.uuid4(), run_dir)


def test_deliverable_bar_accepts_produced_run_end_to_end(tmp_path: Path) -> None:
    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    _produce(run_dir, _build_inputs(run_dir))
    runner._assert_completed_workbook_deliverable(uuid.uuid4(), run_dir)


# ---------------------------------------------------------------------------
# P7 — legacy honesty: the 07W.3 disk read is PROFILE-AGNOSTIC
# ---------------------------------------------------------------------------


def test_p7_legacy_profile_with_activated_contribution_projects_enriched(
    tmp_path: Path,
) -> None:
    """A legacy-profile run dir (no validated brief) WITH an activated 07W.3
    contribution must project the REAL enriched authority — never the
    factually false "no activated workbook-review contribution exists"."""
    from app.marcus.lesson_plan.glossary_projection import (
        render_glossary_projection_markdown,
    )

    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    inputs = wb_act.build_workbook_inputs(
        run_dir, run_id="glos39", validated_brief=None
    )
    assert inputs is not None
    assert inputs.render_profile == "legacy"
    assert inputs.glossary is not None
    assert inputs.glossary.authority == "enriched"
    assert inputs.glossary_empty_reason is None
    rendered = render_glossary_projection_markdown(inputs.glossary)
    assert "no activated workbook-review contribution exists" not in rendered
    assert "### AI" in rendered


# ---------------------------------------------------------------------------
# P6 — B5 supplements derive ONLY from gate-proven content
# ---------------------------------------------------------------------------


def _enriched_payload_with_unused_titled_row(title: str) -> dict[str, Any]:
    from tests.helpers.deep_dive_enrichment_37_2b import live_pool_entry

    request = make_request(
        pool_rows=(
            live_pool_entry(),
            mutated_pool_entry(citation_id="ask-a-cite-002", title=title),
        )
    )
    result = compose_deep_dive_enrichment(request, lambda _: enriched_candidate(request))
    contribution = build_workbook_review_contribution(result, _receipt(request))
    return json.loads(contribution.model_dump_json())


def test_p6_unused_row_title_numeral_never_enters_supplement_set(
    tmp_path: Path,
) -> None:
    """Pin: an UNUSED pool row's title numeral does NOT enter the G1 supplement
    set (titles are not gate-proven) — and the render therefore fails G1 loud
    instead of silently self-licensing the numeral."""
    from app.marcus.lesson_plan.workbook_producer import WorkbookFidelityError

    payload = _enriched_payload_with_unused_titled_row(
        "Healthcare spend reaches $7 trillion in the projection window"
    )
    run_dir = _make_run_dir(tmp_path, payload)
    inputs = _build_inputs(run_dir)
    # The prose-unused ask-a-cite-002 row association-covers "AI" (rendered in
    # the glossary), but its TITLE numeral is NOT declared as a supplement.
    assert "money-trillion:7" not in inputs.research_supplements
    with pytest.raises(WorkbookFidelityError, match="unsourced"):
        _produce(run_dir, inputs)


def test_p6_used_row_evidence_excerpt_figures_are_declared(tmp_path: Path) -> None:
    """The USED row's gate-checked evidence excerpt stays a declared supplement
    (the tightening drops titles/unused rows, not gate-proven content)."""
    from app.specialists._shared.figure_tokens import _figures

    run_dir = _make_run_dir(tmp_path, _enriched_contribution_payload())
    inputs = _build_inputs(run_dir)
    result = inputs.deep_dive_review.deep_dive_enrichment
    from app.marcus.lesson_plan.deep_dive_enrichment import used_pool_rows

    excerpt_figures = set()
    for row in used_pool_rows(result):
        excerpt_figures |= _figures(row.evidence_excerpt)
    assert excerpt_figures <= inputs.research_supplements
